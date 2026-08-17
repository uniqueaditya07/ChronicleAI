import docx
from app.services.parsers.docx_parser import DOCXParser
from app.models.source import ChunkType

def test_docx_parser_can_handle():
    parser = DOCXParser()
    assert parser.can_handle("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx")
    assert not parser.can_handle("text/plain", "txt")

def test_docx_parser_parsing(tmp_path):
    # Create a docx file for testing
    doc = docx.Document()
    doc.add_paragraph("This is a heading", style='Heading 1')
    doc.add_paragraph("This is a paragraph.")
    
    table = doc.add_table(rows=1, cols=2)
    row_cells = table.rows[0].cells
    row_cells[0].text = "Cell 1"
    row_cells[1].text = "Cell 2"
    
    test_file = tmp_path / "test.docx"
    doc.save(test_file)
    
    parser = DOCXParser()
    parsed_doc = parser.parse(test_file)
    
    assert len(parsed_doc.chunks) == 3
    
    # Paragraph 1
    assert parsed_doc.chunks[0].chunk_type == ChunkType.paragraph
    assert parsed_doc.chunks[0].content_text == "This is a heading"
    assert parsed_doc.chunks[0].metadata["docx_paragraph_index"] == 0
    
    # Paragraph 2
    assert parsed_doc.chunks[1].chunk_type == ChunkType.paragraph
    assert parsed_doc.chunks[1].content_text == "This is a paragraph."
    
    # Table Row
    assert parsed_doc.chunks[2].chunk_type == ChunkType.row
    assert parsed_doc.chunks[2].content_text == "Cell 1 | Cell 2"
    assert parsed_doc.chunks[2].metadata["docx_table_index"] == 0
    assert parsed_doc.chunks[2].metadata["docx_row_index"] == 0
